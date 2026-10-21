import csv
import io
import json
from pathlib import Path

from pydantic import ValidationError

from app.common.exceptions import AppException
from app.modules.quiz.schema import QuestionIn

MAX_BYTES, MAX_ROWS = 2 * 1024 * 1024, 2000
ALLOWED_EXT = {".csv", ".json", ".xlsx"}
ALLOWED_MIME = {"text/csv", "application/json", "text/plain", "application/octet-stream",
                "application/vnd.ms-excel",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


def parse_rows(filename: str, content_type: str | None, content: bytes) -> list[dict]:
    ext = Path(filename or "").suffix.lower()  # filename is only inspected, never used as a path
    if ext not in ALLOWED_EXT or (content_type and content_type not in ALLOWED_MIME):
        raise AppException(400, "UNSUPPORTED_FILE", "Only .csv, .json, .xlsx files are allowed")
    if len(content) > MAX_BYTES:
        raise AppException(413, "FILE_TOO_LARGE", "File exceeds 2MB")
    try:
        if ext == ".json":
            rows = json.loads(content.decode("utf-8-sig"))
            if not isinstance(rows, list):
                raise ValueError("JSON root must be a list")
        elif ext == ".csv":
            rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        else:
            from openpyxl import load_workbook

            ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            it = ws.iter_rows(values_only=True)
            header = [str(h or "").strip().lower() for h in next(it, [])]
            rows = [dict(zip(header, r, strict=False)) for r in it if any(c is not None for c in r)]
    except (ValueError, UnicodeDecodeError, KeyError, OSError) as exc:
        raise AppException(400, "PARSE_ERROR", f"Cannot parse file: {exc}") from exc
    if not rows or len(rows) > MAX_ROWS:
        raise AppException(400, "INVALID_ROW_COUNT", f"File must contain 1-{MAX_ROWS} rows")
    return [r for r in rows if isinstance(r, dict)]


def normalize_row(row: dict) -> dict:
    r = {str(k).strip().lower(): (v.strip() if isinstance(v, str) else v)
         for k, v in row.items() if k is not None}
    if "options" not in r:
        opts = [str(r[f"option_{c}"]) for c in "abcd" if r.get(f"option_{c}") not in (None, "")]
        if opts:
            r["options"] = opts
    ans = r.pop("answer", None)
    if "correct_answer" not in r and ans is not None:
        a, opts = str(ans).strip(), r.get("options") or []
        if len(a) == 1 and a.upper() in "ABCD" and ord(a.upper()) - 65 < len(opts):
            r["correct_answer"] = opts[ord(a.upper()) - 65]
        else:
            r["correct_answer"] = a
    for c in "abcd":
        r.pop(f"option_{c}", None)
    r.setdefault("type", "multiple_choice")
    return r


def validate_rows(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    valid, errors = [], []
    for i, row in enumerate(rows, start=1):
        try:
            valid.append(QuestionIn(**normalize_row(row)).model_dump())
        except (ValidationError, TypeError) as exc:
            msg = exc.errors()[0]["msg"] if isinstance(exc, ValidationError) else str(exc)
            errors.append({"row": i, "error": msg})
    return valid, errors[:50]
